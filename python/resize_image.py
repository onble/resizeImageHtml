#!/usr/bin/env python3
"""Resize image assets to a fixed height while preserving folders."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path


TOOL_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = TOOL_DIR / "input"
DEFAULT_OUTPUT = TOOL_DIR / "output"
SUPPORTED_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
DEFAULT_SIZE_MAP_SHEET = "怪物属性表"
MONSTER_ID_PATTERN = re.compile(r"mon(\d+)", re.IGNORECASE)


def positive_int(value: str) -> int:
    try:
        number = int(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError(f"{value!r} is not an integer") from exc
    if number <= 0:
        raise argparse.ArgumentTypeError("size must be greater than 0")
    return number


def load_pillow():
    try:
        from PIL import Image
    except ImportError:
        print(
            "Pillow is required. Install it with:\n"
            "  py -3 -m pip install Pillow\n"
            "or:\n"
            "  python -m pip install Pillow",
            file=sys.stderr,
        )
        return None
    return Image


def load_openpyxl():
    try:
        from openpyxl import load_workbook
    except ImportError:
        print(
            "openpyxl is required when using --size-map. Install it with:\n"
            "  py -3 -m pip install openpyxl\n"
            "or:\n"
            "  python -m pip install openpyxl",
            file=sys.stderr,
        )
        return None
    return load_workbook


def image_files(input_path: Path) -> list[Path]:
    if input_path.is_file():
        if input_path.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise ValueError(f"input file is not a supported image: {input_path}")
        return [input_path]
    if input_path.is_dir():
        return sorted(
            path
            for path in input_path.rglob("*")
            if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
        )
    raise FileNotFoundError(f"input path does not exist: {input_path}")


def parse_size(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float):
        return int(value) if value > 0 and value.is_integer() else None

    match = re.search(r"\d+", str(value))
    if not match:
        return None
    size = int(match.group(0))
    return size if size > 0 else None


def load_size_map(mapping_path: Path, sheet_name: str) -> dict[int, int]:
    load_workbook = load_openpyxl()
    if load_workbook is None:
        raise RuntimeError("openpyxl is not available")

    workbook = load_workbook(mapping_path, data_only=True, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"sheet not found in size map: {sheet_name}")

        sheet = workbook[sheet_name]
        header_row = None
        id_col = None
        size_col = None

        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10)):
            values = [cell.value for cell in row]
            if values and str(values[0]).strip().upper() == "ID":
                for index, value in enumerate(values, start=1):
                    if value is not None and "帧动画尺寸" in str(value):
                        header_row = row[0].row
                        id_col = 1
                        size_col = index
                        break
            if header_row is not None:
                break

        if header_row is None or id_col is None or size_col is None:
            header_row = 2
            id_col = 1
            size_col = 8

        size_map: dict[int, int] = {}
        for row_index in range(header_row + 1, sheet.max_row + 1):
            raw_id = sheet.cell(row_index, id_col).value
            raw_size = sheet.cell(row_index, size_col).value
            try:
                monster_id = int(raw_id)
            except (TypeError, ValueError):
                continue

            size = parse_size(raw_size)
            if size is None:
                continue
            size_map[monster_id] = size
    finally:
        workbook.close()

    if not size_map:
        raise ValueError(f"no usable ID/size rows found in size map: {mapping_path}")
    return size_map


def monster_id_for(source: Path, input_path: Path) -> int | None:
    try:
        parts = source.relative_to(input_path).parts if input_path.is_dir() else source.parts
    except ValueError:
        parts = source.parts

    for part in parts:
        match = MONSTER_ID_PATTERN.search(part)
        if match:
            return int(match.group(1))
    return None


def output_path_for(source: Path, input_path: Path, output_path: Path) -> Path:
    if input_path.is_file():
        if output_path.suffix.lower() == ".png":
            return output_path
        return output_path / source.name
    return output_path / source.relative_to(input_path)


def resize_png(
    source: Path,
    destination: Path,
    target_height: int,
    resample,
    image_module,
    palette: bool,
    square: bool,
) -> tuple[int, tuple[int, int], tuple[int, int], int]:
    with image_module.open(source) as image:
        image.load()
        original_size = image.size

        # Keep alpha clean for UI/game sprites and strip metadata on save.
        working = image.convert("RGBA")
        original_width, original_height = original_size
        target_width = (
            target_height
            if square
            else max(1, round(original_width * target_height / original_height))
        )
        target_size = (target_width, target_height)
        resized = working.resize(target_size, resample)

        destination.parent.mkdir(parents=True, exist_ok=True)
        if destination.suffix.lower() == ".png" and palette:
            output_image = resized.convert("P", palette=image_module.Palette.ADAPTIVE, colors=256)
            output_image.save(destination, format="PNG", optimize=True, compress_level=9)
        elif destination.suffix.lower() == ".png":
            resized.save(destination, format="PNG", optimize=True, compress_level=9)
        else:
            resized.convert("RGB").save(destination, optimize=True, quality=95)

    return source.stat().st_size, original_size, target_size, destination.stat().st_size


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Resize image files for game assets. Defaults to ./input -> ./output."
    )
    parser.add_argument(
        "size",
        nargs="?",
        type=positive_int,
        default=120,
        help="target image height in pixels; width is scaled proportionally",
    )
    parser.add_argument(
        "-i",
        "--input",
        type=Path,
        default=DEFAULT_INPUT,
        help="input image file or directory; defaults to tools/resizeImage/input",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="output image file or directory; defaults to tools/resizeImage/output",
    )
    parser.add_argument(
        "--nearest",
        action="store_true",
        help="use nearest-neighbor scaling for pixel-art assets",
    )
    parser.add_argument(
        "--palette",
        action="store_true",
        help="convert output PNG files to a 256-color palette for smaller files",
    )
    parser.add_argument(
        "--square",
        action="store_true",
        help="resize output to size x size instead of preserving the aspect ratio",
    )
    parser.add_argument(
        "--size-map",
        type=Path,
        help="Excel mapping file whose first ID column maps monsters to per-frame output sizes",
    )
    parser.add_argument(
        "--size-map-sheet",
        default=DEFAULT_SIZE_MAP_SHEET,
        help=f"worksheet name for --size-map; defaults to {DEFAULT_SIZE_MAP_SHEET}",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    image_module = load_pillow()
    if image_module is None:
        return 1

    input_path = args.input.resolve()
    output_path = args.output.resolve()
    size_map = None

    if args.size_map:
        try:
            size_map = load_size_map(args.size_map.resolve(), args.size_map_sheet)
        except (OSError, ValueError, RuntimeError) as exc:
            print(f"Error: {exc}", file=sys.stderr)
            return 1

    try:
        files = image_files(input_path)
    except (FileNotFoundError, ValueError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    if not files:
        print(f"No supported image files found in: {input_path}")
        return 0

    resample = image_module.Resampling.NEAREST if args.nearest else image_module.Resampling.LANCZOS

    print(f"Input : {input_path}")
    print(f"Output: {output_path}")
    if args.square:
        print(f"Size  : {args.size}x{args.size}px")
    elif size_map is None:
        print(f"Height: {args.size}px, width scaled proportionally")
    else:
        print(f"Height: mapped per monster from {args.size_map.resolve()}")
        print(f"Map   : {len(size_map)} monster size entr{'y' if len(size_map) == 1 else 'ies'}")
    print(f"Mode  : {'NEAREST' if args.nearest else 'LANCZOS'}")
    if args.palette:
        print("PNG   : optimized palette, 256 colors")
    else:
        print("PNG   : optimized RGBA")
    print()

    converted = 0
    for source in files:
        destination = output_path_for(source, input_path, output_path)
        relative_name = source.name if input_path.is_file() else str(source.relative_to(input_path))
        target_size = args.size
        if size_map is not None:
            monster_id = monster_id_for(source, input_path)
            if monster_id is None:
                print(f"FAIL {relative_name}: could not find monster ID like mon01 in path", file=sys.stderr)
                continue
            target_size = size_map.get(monster_id)
            if target_size is None:
                print(f"FAIL {relative_name}: no output size mapped for monster ID {monster_id}", file=sys.stderr)
                continue

        try:
            before_bytes, original_size, resized_size, after_bytes = resize_png(
                source,
                destination,
                target_size,
                resample,
                image_module,
                args.palette,
                args.square,
            )
        except Exception as exc:  # Keep batch conversion going.
            print(f"FAIL {source}: {exc}", file=sys.stderr)
            continue

        converted += 1
        before_kb = before_bytes / 1024
        after_kb = after_bytes / 1024
        print(
            f"OK   {relative_name}: {original_size[0]}x{original_size[1]} -> "
            f"{resized_size[0]}x{resized_size[1]}, {before_kb:.1f} KB -> {after_kb:.1f} KB"
        )

    print()
    print(f"Converted {converted}/{len(files)} image file(s).")
    return 0 if converted == len(files) else 1


if __name__ == "__main__":
    raise SystemExit(main())
